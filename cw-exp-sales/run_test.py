"""サンプル集計と、読み取り専用・エラー時の動作を確認します。"""

import csv
from contextlib import redirect_stdout
import hashlib
import io
import os
from pathlib import Path
import shutil
import subprocess
import sys
import tempfile
from unittest.mock import patch

from openpyxl import load_workbook

import main as sales_tool


BASE_DIR = Path(__file__).resolve().parent
HEADERS = ["日付", "店舗名", "商品名", "数量", "単価"]
SAMPLE_ROWS = [
    ("2026/10/01", "草津店", "商品A", 2, 500, 1000),
    ("2026/10/01", "大津店", "商品B", 1, 1200, 1200),
    ("2026/10/02", "草津店", "商品B", 3, 1200, 3600),
    ("2026/10/03", "大津店", "商品A", 4, 500, 2000),
    ("2026/10/03", "彦根店", "商品C", 2, 800, 1600),
    ("2026/10/04", "草津店", "商品C", 1, 800, 800),
    ("2026/10/05", "大津店", "商品B", 2, 1200, 2400),
    ("2026/10/05", "彦根店", "商品A", 3, 500, 1500),
    ("2026/10/06", "草津店", "商品A", 1, 500, 500),
]
SAMPLE_TOTALS = {"草津店": 5900, "大津店": 5600, "彦根店": 3100}


def hashes(folder):
    # バイト列のハッシュで、文字コードや改行も含めた不変性を確認します。
    return {
        path.name: hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(folder.glob("*.csv"))
    }


def run_tool(script, arguments, cwd, expected_code=0):
    environment = os.environ.copy()
    environment.pop("CW_EXP_SALES_PAUSE", None)
    environment["PYTHONUTF8"] = "1"
    result = subprocess.run(
        [sys.executable, str(script), *map(str, arguments)],
        cwd=cwd,
        env=environment,
        stdin=subprocess.DEVNULL,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=60,
    )
    assert result.returncode == expected_code, result.stdout + result.stderr
    assert not result.stderr, result.stderr
    assert "Enterキーを押すと終了します。" not in result.stdout
    return result.stdout


def check_workbook(path, expected_rows, expected_totals):
    with path.open("rb") as source:
        workbook = load_workbook(source, data_only=False)
        try:
            assert workbook.sheetnames == ["明細", "店舗別合計"]
            details = list(workbook["明細"].values)
            assert details[0] == tuple(HEADERS + ["売上金額"])
            assert details[1:] == expected_rows, details
            totals = list(workbook["店舗別合計"].values)
            assert totals[0] == ("店舗名", "売上合計")
            assert len(totals) == len(expected_totals) + 1
            assert dict(totals[1:]) == expected_totals, totals
            assert sum(row[5] for row in details[1:]) == sum(expected_totals.values())
            for row in workbook["明細"].iter_rows(min_row=2):
                assert all(cell.data_type == "s" for cell in row[:3])
                assert all(cell.data_type == "n" for cell in row[3:])
        finally:
            workbook.close()


def write_csv(path, rows, encoding="utf-8-sig"):
    # 検証用のCSVだけを一時フォルダーに作成します。
    with path.open("w", encoding=encoding, newline="") as target:
        csv.writer(target).writerows(rows)


def main():
    sample_dir = BASE_DIR / "input_csv"
    before = hashes(sample_dir)
    assert set(before) == {"01_sales_utf8.csv", "02_sales_utf8.csv", "03_sales_cp932.csv"}
    for filename in ("01_sales_utf8.csv", "02_sales_utf8.csv"):
        data = (sample_dir / filename).read_bytes()
        assert data.startswith(b"\xef\xbb\xbf")
        assert next(csv.reader(io.StringIO(data.decode("utf-8-sig")))) == HEADERS
    cp932_data = (sample_dir / "03_sales_cp932.csv").read_bytes()
    assert not cp932_data.startswith(b"\xef\xbb\xbf")
    assert next(csv.reader(io.StringIO(cp932_data.decode("cp932"))))[0] == "2026/10/05"
    try:
        cp932_data.decode("utf-8-sig")
    except UnicodeDecodeError:
        pass
    else:
        raise AssertionError("CP932のサンプルで文字コードの再試行を検証できません。")

    with tempfile.TemporaryDirectory(prefix="cw-exp-sales-test-") as temporary:
        temp_dir = Path(temporary)
        other_cwd = temp_dir / "別の起動場所"
        other_cwd.mkdir()

        # 引数なしでも、起動場所によらず本物のサンプルと出力先を使います。
        stdout = run_tool(BASE_DIR / "main.py", [], other_cwd)
        print(stdout, end="")
        check_workbook(BASE_DIR / "output" / "売上集計.xlsx", SAMPLE_ROWS, SAMPLE_TOTALS)
        assert "【警告】" not in stdout
        assert before == hashes(sample_dir), "サンプルCSVが変更・削除されています。"
        print("[PASS] サンプル3ファイル、9明細、店舗別合計、全店合計14,600円")
        print("[PASS] サンプルCSVのSHA-256は全3ファイルで実行前後一致")

        # 異常系はスクリプトも一時コピーし、同梱Excelに影響させません。
        tool_dir = temp_dir / "検証ツール"
        tool_dir.mkdir()
        script = tool_dir / "main.py"
        shutil.copyfile(BASE_DIR / "main.py", script)
        fixture_dir = tool_dir / "検証 CSV"
        fixture_dir.mkdir()
        write_csv(fixture_dir / "01_mixed.csv", [
            [],
            [" 日付 ", "店舗名", "商品名", "数量", "単価"],
            ["2026/10/01", "検証店", "=商品A", 2, 500],
            ["2026/10/01", "検証店", "商品A", "二", 500],
            ["2026/10/01", "検証店", "商品A", 1, "不明"],
            ["2026/10/01", "検証店", "商品A", "NaN", 500],
            ["2026/10/01", "検証店", "商品A", 1, "inf"],
            ["2026/10/01", "検証店", "商品A", "", 500],
            ["2026/10/01", "検証店", "商品A", "1e308", "1e308"],
            ["2026/10/01", "検証店", "商品A", 2],
            ["2026/10/01", "", "商品A", 1, 500],
            ["2026/10/01", "検証店", "商品C", 1.5, 200],
            ["2026/10/01", "検証店", "商品A", -1, 500],
        ])
        write_csv(fixture_dir / "02_cp932.csv", [
            ["2026/10/02", "別店舗", "商品B", 1, 1200],
        ], encoding="cp932")
        write_csv(fixture_dir / "03_empty.csv", [])
        write_csv(fixture_dir / "04_header_only.csv", [HEADERS])
        (fixture_dir / "05_unreadable.csv").write_bytes(b"\x81")
        (fixture_dir / "06_malformed.csv").write_text('"閉じていない引用符', encoding="utf-8")
        (fixture_dir / "07_blank.csv").write_text("\n  \n,,,\n", encoding="utf-8-sig")
        # CSV以外とサブフォルダーは対象外です。
        (fixture_dir / "ignored.txt").write_text("集計しません", encoding="utf-8")
        nested = fixture_dir / "サブフォルダー"
        nested.mkdir()
        write_csv(nested / "ignored.csv", [["2026/10/01", "除外店", "商品A", 1, 99999]])
        expected_rows = [
            ("2026/10/01", "検証店", "=商品A", 2, 500, 1000),
            ("2026/10/01", "検証店", "商品C", 1.5, 200, 300),
            ("2026/10/01", "検証店", "商品A", -1, 500, -500),
            ("2026/10/02", "別店舗", "商品B", 1, 1200, 1200),
        ]
        expected_totals = {"検証店": 800, "別店舗": 1200}
        fixture_before = hashes(fixture_dir)
        output_path = tool_dir / "output" / "売上集計.xlsx"
        for argument in (Path("検証 CSV"), fixture_dir):
            stdout = run_tool(script, [argument], other_cwd)
            check_workbook(output_path, expected_rows, expected_totals)
            assert fixture_before == hashes(fixture_dir)
            for filename in ("03_empty.csv", "04_header_only.csv", "05_unreadable.csv",
                             "06_malformed.csv", "07_blank.csv"):
                assert any("【警告】" in line and filename in line for line in stdout.splitlines())
            assert stdout.count("正しい数値ではない") == 6, stdout
            assert "01_mixed.csv の 4 行目" in stdout
            assert "列数が5列ではない" in stdout
            assert "空欄がある" in stdout
        print("[PASS] UTF-8/CP932、見出し自動判定、相対/絶対パス、別の起動場所")
        print("[PASS] 空・読取不能・不正CSV、不正数値、列数不正でも有効な行を継続集計")
        print("[PASS] 小数・負数、文字列の保持、対象外ファイルとサブフォルダーの除外")
        print("[PASS] 異常系CSVもSHA-256が実行前後一致")

        # OSのアクセス拒否も、元ファイルを変更せずに再現します。
        warning = io.StringIO()
        with patch.object(Path, "open", side_effect=PermissionError) as opened:
            with redirect_stdout(warning):
                assert sales_tool.read_sales(sample_dir / "01_sales_utf8.csv") == []
            opened.assert_called_once_with("r", encoding="utf-8-sig", newline="")
        assert "【警告】" in warning.getvalue() and "読み取れない" in warning.getvalue()
        print("[PASS] アクセス拒否は日本語警告でスキップし、CSVを読み取り専用で開く")

        # 有効な行がないときには、既存Excelを上書きしないことも確認します。
        output_hash = hashlib.sha256(output_path.read_bytes()).hexdigest()
        invalid_dir = tool_dir / "無効なCSVだけ"
        invalid_dir.mkdir()
        write_csv(invalid_dir / "invalid.csv", [["2026/10/01", "検証店", "商品A", "不明", 500]])
        stdout = run_tool(script, [invalid_dir], other_cwd, expected_code=1)
        assert "Excelは作成・更新していません" in stdout
        empty_dir = tool_dir / "CSVなし"
        empty_dir.mkdir()
        stdout = run_tool(script, [empty_dir], other_cwd, expected_code=1)
        assert "CSVファイルがありません" in stdout
        stdout = run_tool(script, [tool_dir / "存在しない"], other_cwd, expected_code=1)
        assert "CSVフォルダーが見つかりません" in stdout
        stdout = run_tool(script, [fixture_dir, "余分な引数"], other_cwd, expected_code=1)
        assert "使い方:" in stdout
        assert output_hash == hashlib.sha256(output_path.read_bytes()).hexdigest()
        print("[PASS] 有効行なし・CSVなし・フォルダーなし・引数過多: 終了コード1、既存Excelを維持")

    assert before == hashes(sample_dir)
    print("テストはすべて成功しました。元CSVは変更されていません。")
    print(f"サンプルExcel: {BASE_DIR / 'output' / '売上集計.xlsx'}")


if __name__ == "__main__":
    main()
