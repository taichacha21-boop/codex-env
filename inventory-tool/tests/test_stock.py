from __future__ import annotations

import csv
import hashlib
import io
import os
import shutil
import tempfile
import unittest
from contextlib import closing
from copy import copy
from datetime import datetime
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook

import main


SAMPLES = Path(__file__).resolve().parents[1] / "sample_data"
HEADERS = ["伝票番号", "明細ID", "日付", "商品コード", "区分", "数量"]


def record(slip="T001", detail="001", day="2026-10-01", code="00125", kind="入庫", quantity="5"):
    return [slip, detail, day, code, kind, quantity]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


class FixedDateTime(datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 10, 27, 15, 30, 0, tzinfo=tz)


class StockTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.base = Path(self.temporary.name)
        self.master = self.base / "master.xlsx"
        shutil.copyfile(SAMPLES / "master.xlsx", self.master)
        self.sequence = 0

    def folder(self, rows=None, headers=HEADERS, encoding="utf-8-sig"):
        self.sequence += 1
        folder = self.base / f"csv_{self.sequence}"
        folder.mkdir()
        if rows is not None:
            self.write_csv(folder / "data.csv", rows, headers, encoding)
        return folder

    def write_csv(self, path, rows, headers=HEADERS, encoding="utf-8-sig"):
        with path.open("w", encoding=encoding, newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(rows)

    def sample_folder(self, name):
        folder = self.folder()
        for path in (SAMPLES / name).iterdir():
            shutil.copyfile(path, folder / path.name)
        return folder

    def edit_master(self, edit, path=None):
        path = path or self.master
        with closing(load_workbook(path)) as wb:
            edit(wb)
            wb.save(path)

    def stocks(self, path):
        with closing(load_workbook(path, data_only=False)) as wb:
            return {str(row[0].value): row[2].value for row in wb["商品マスタ"].iter_rows(min_row=2)}

    def history(self, path):
        with closing(load_workbook(path)) as wb:
            return list(wb["処理履歴"].values)

    def assert_success(self, result):
        self.assertTrue(result.success, result.errors)
        self.assertTrue(result.output_path.is_file())

    def assert_stopped(self, result, message=None):
        self.assertFalse(result.success)
        self.assertIsNone(result.output_path)
        self.assertEqual(result.applied_count, 0)
        self.assertTrue(result.errors)
        if message:
            self.assertTrue(any(message in issue.message for issue in result.errors), result.errors)

    def outputs(self):
        return set(self.base.glob("商品マスタ_更新結果_*.xlsx"))

    def test_normal_multiple_files_both_encodings_and_leading_zero_codes(self):
        folder = self.sample_folder("normal")
        result = main.update_stock(self.master, folder, "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 4)
        self.assertEqual(result.errors, [])
        self.assertEqual(result.warnings, [])
        self.assertEqual(self.stocks(result.output_path), {"00125": 107, "00408": 18, "00999": 2})
        with closing(load_workbook(result.output_path)) as wb:
            summary = wb[result.report_sheets["集計"]]
            self.assertEqual([c.value for c in summary[1]], list(main.SUMMARY_HEADERS))
            self.assertEqual([c.value for c in summary[2]][:6], ["00125", "商品A", 100, 10, 3, 107])
            for sheet, column in ((wb["商品マスタ"], 1), (summary, 1), (wb["処理履歴"], 4)):
                for row in sheet.iter_rows(min_row=2):
                    self.assertEqual(row[column - 1].data_type, "s")
            self.assertEqual(wb["エラー一覧"].max_row, 1)
            self.assertEqual(wb["警告一覧"].max_row, 1)
        self.assertEqual(len(self.history(result.output_path)), 5)

    def test_identical_keys_within_multiple_files_apply_once(self):
        result = main.update_stock(self.master, self.sample_folder("duplicate"), "first")
        self.assert_success(result)
        self.assertEqual((result.applied_count, result.duplicate_count), (2, 2))
        self.assertEqual(self.stocks(result.output_path)["00125"], 110)
        self.assertEqual(len(self.history(result.output_path)), 3)

    def test_continue_same_csv_is_idempotent(self):
        folder = self.sample_folder("normal")
        first = main.update_stock(self.master, folder, "first")
        self.assert_success(first)
        before = digest(first.output_path)
        second = main.update_stock(first.output_path, folder, "continue")
        self.assert_success(second)
        self.assertEqual((second.applied_count, second.duplicate_count), (0, 4))
        self.assertEqual(self.stocks(first.output_path), self.stocks(second.output_path))
        self.assertEqual(self.history(first.output_path), self.history(second.output_path))
        self.assertEqual(before, digest(first.output_path))
        self.assertEqual(second.report_sheets["集計"], "集計_2")

    def test_continue_mixed_old_and_new_keys(self):
        original = record()
        first = main.update_stock(self.master, self.folder([original]), "first")
        second = main.update_stock(first.output_path, self.folder([original, record(slip="NEW", kind="出庫", quantity="8")]), "continue")
        self.assert_success(second)
        self.assertEqual((second.applied_count, second.duplicate_count), (1, 1))
        self.assertEqual(self.stocks(second.output_path)["00125"], 97)
        self.assertEqual(len(self.history(second.output_path)), 3)

    def test_each_conflicting_field_aborts_entire_run(self):
        for field, changed in ((2, "2026-10-02"), (3, "00408"), (4, "出庫"), (5, "6")):
            with self.subTest(field=HEADERS[field]):
                folder = self.folder([record(), record(slip="GOOD", quantity="12")])
                conflicting = record()
                conflicting[field] = changed
                self.write_csv(folder / "later.csv", [conflicting])
                before = digest(self.master)
                result = main.update_stock(self.master, folder, "first")
                self.assert_stopped(result, "矛盾")
                self.assertEqual(self.outputs(), set())
                self.assertEqual(digest(self.master), before)
                conflict = next(issue for issue in result.errors if "矛盾" in issue.message)
                self.assertEqual((conflict.source_file, conflict.row), ("later.csv", 2))

    def test_conflicting_history_field_aborts_instead_of_partially_applying(self):
        first = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(first)
        previous_hash = digest(first.output_path)
        for field, changed in ((2, "2026-10-02"), (3, "00408"), (4, "出庫"), (5, "9")):
            with self.subTest(field=HEADERS[field]):
                row = record()
                row[field] = changed
                result = main.update_stock(first.output_path, self.folder([record(slip="GOOD"), row]), "continue")
                self.assert_stopped(result, "矛盾")
                self.assertEqual(self.outputs(), {first.output_path})
                self.assertEqual(digest(first.output_path), previous_hash)

    def test_invalid_or_unknown_rows_still_participate_in_conflict_detection(self):
        for change in (record(quantity="0"), record(code="UNKNOWN"), record(day="invalid"), record(kind="返品")):
            with self.subTest(change=change):
                result = main.update_stock(self.master, self.folder([record(), change]), "first")
                self.assert_stopped(result, "矛盾")
                self.assertEqual(self.outputs(), set())

    def test_row_errors_are_separate_from_warnings_and_do_not_enter_history(self):
        result = main.update_stock(self.master, self.sample_folder("errors"), "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        self.assertEqual(len(result.errors), 10)
        self.assertEqual(result.warnings, [])
        self.assertEqual(self.stocks(result.output_path)["00125"], 102)
        self.assertEqual(self.stocks(result.output_path)["00408"], 20)
        self.assertEqual(len(self.history(result.output_path)), 2)
        self.assertTrue(any("408" in issue.message and "未登録" in issue.message for issue in result.errors))
        self.assertEqual([issue.row for issue in result.errors], list(range(2, 12)))
        self.assertTrue(all(issue.source_file == "errors.csv" for issue in result.errors))
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["エラー一覧"].max_row, 11)
            self.assertEqual(wb["警告一覧"].max_row, 1)

    def test_repaired_error_row_can_be_imported_later(self):
        first = main.update_stock(self.master, self.folder([record(detail=""), record(slip="GOOD")]), "first")
        second = main.update_stock(first.output_path, self.folder([record(), record(slip="GOOD")]), "continue")
        self.assert_success(second)
        self.assertEqual((second.applied_count, second.duplicate_count), (1, 1))
        self.assertEqual(self.stocks(second.output_path)["00125"], 110)

    def test_duplicate_master_code_is_fatal_after_trim(self):
        shutil.copyfile(SAMPLES / "master_duplicate.xlsx", self.master)
        self.edit_master(lambda wb: setattr(wb["商品マスタ"]["A5"], "value", " 00125 "))
        before = digest(self.master)
        result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_stopped(result, "商品コード重複")
        self.assertEqual(self.outputs(), set())
        self.assertEqual(digest(self.master), before)
        self.assertEqual(result.errors[0].row, 5)

    def test_numeric_master_code_is_string_without_invented_leading_zero(self):
        def change(wb):
            wb["商品マスタ"]["A3"] = 408
            wb["商品マスタ"]["A3"].number_format = "00000"
        self.edit_master(change)
        result = main.update_stock(self.master, self.folder([record(code="408"), record(slip="OTHER", code="00408")]), "first")
        self.assert_success(result)
        self.assertEqual(self.stocks(result.output_path)["408"], 25)
        self.assertTrue(any("00408" in issue.message for issue in result.errors))
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["商品マスタ"]["A3"].value, "408")
            self.assertEqual(wb["商品マスタ"]["A3"].data_type, "s")
            self.assertEqual(wb["商品マスタ"]["A3"].number_format, "00000")
            self.assertEqual(wb["処理履歴"]["D2"].value, "408")

    def test_formula_stock_skips_update_warns_and_is_not_recorded(self):
        shutil.copyfile(SAMPLES / "master_formula.xlsx", self.master)
        folder = self.sample_folder("warnings")
        first = main.update_stock(self.master, folder, "first")
        self.assert_success(first)
        self.assertEqual(first.applied_count, 1)
        self.assertEqual(first.errors, [])
        self.assertEqual(self.stocks(first.output_path)["F001"], "=SUM(C2:C3)")
        self.assertEqual(self.stocks(first.output_path)["00999"], -3)
        self.assertEqual(len(self.history(first.output_path)), 2)
        self.assertEqual([(issue.source_file, issue.row) for issue in first.warnings], [("warnings.csv", 2), ("warnings.csv", 3)])
        with closing(load_workbook(first.output_path)) as wb:
            summary = wb["集計"]
            self.assertEqual(summary["G5"].value, "現在庫が数式のため未更新")
            self.assertEqual(summary["D5"].value, 0)
            self.assertEqual(summary["F5"].value, "=SUM(C2:C3)")
            self.assertEqual(summary["F5"].data_type, "s")
            self.assertEqual(wb["商品マスタ"]["C5"].data_type, "f")
        repaired = self.base / "repaired.xlsx"
        shutil.copyfile(first.output_path, repaired)
        self.edit_master(lambda wb: setattr(wb["商品マスタ"]["C5"], "value", 20), repaired)
        second = main.update_stock(repaired, folder, "continue")
        self.assert_success(second)
        self.assertEqual((second.applied_count, second.duplicate_count), (1, 1))
        self.assertEqual(self.stocks(second.output_path)["F001"], 24)
        self.assertEqual(self.stocks(second.output_path)["00999"], -3)

    def test_negative_stock_remains_negative_and_warns_with_source_row(self):
        result = main.update_stock(self.master, self.folder([record(code="00999", kind="出庫", quantity="8")]), "first")
        self.assert_success(result)
        self.assertEqual(self.stocks(result.output_path)["00999"], -6)
        self.assertEqual(result.errors, [])
        self.assertEqual(len(result.warnings), 1)
        self.assertEqual((result.warnings[0].source_file, result.warnings[0].row), ("data.csv", 2))
        self.assertIn("負在庫", result.warnings[0].message)
        self.assertEqual(len(self.history(result.output_path)), 2)

    def test_invalid_master_row_is_excluded_while_other_product_updates(self):
        self.edit_master(lambda wb: setattr(wb["商品マスタ"]["C2"], "value", "100"))
        result = main.update_stock(self.master, self.folder([record(), record(slip="OTHER", code="00408")]), "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        self.assertEqual(self.stocks(result.output_path)["00125"], "100")
        self.assertEqual(self.stocks(result.output_path)["00408"], 25)
        self.assertEqual(len(result.errors), 2)
        self.assertEqual(len(self.history(result.output_path)), 2)

    def test_master_sheet_and_required_headers_are_mandatory(self):
        for case in ("sheet", "商品コード", "商品名", "現在庫", "duplicate_header"):
            with self.subTest(case=case):
                shutil.copyfile(SAMPLES / "master.xlsx", self.master)
                def change(wb):
                    ws = wb["商品マスタ"]
                    if case == "sheet":
                        ws.title = "別名"
                    elif case == "duplicate_header":
                        ws["F1"] = "現在庫"
                    else:
                        for cell in ws[1]:
                            if cell.value == case:
                                cell.value = "不足"
                self.edit_master(change)
                result = main.update_stock(self.master, self.folder([record()]), "first")
                self.assert_stopped(result)
                self.assertEqual(self.outputs(), set())

    def test_unit_price_optional_and_extra_columns_retained(self):
        self.edit_master(lambda wb: setattr(wb["商品マスタ"]["D1"], "value", "任意列"))
        result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(result)
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["商品マスタ"]["D1"].value, "任意列")
            self.assertEqual(wb["商品マスタ"]["D2"].value, 120)

    def test_existing_product_name_formula_is_preserved_without_blocking_stock_update(self):
        self.edit_master(lambda wb: setattr(wb["商品マスタ"]["B2"], "value", '=CONCAT("商品", "A")'))
        result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["商品マスタ"]["B2"].value, '=CONCAT("商品", "A")')
            self.assertEqual(wb["商品マスタ"]["B2"].data_type, "f")
            self.assertEqual(wb["商品マスタ"]["C2"].value, 105)
            self.assertEqual(wb["集計"]["B2"].data_type, "s")

    def test_mode_must_be_explicit_and_continue_never_falls_back(self):
        for mode in ("", None, "auto", "continue"):
            with self.subTest(mode=mode):
                result = main.update_stock(self.master, self.folder([record()]), mode)
                self.assert_stopped(result)
                self.assertEqual(self.outputs(), set())

    def test_first_mode_cannot_discard_existing_history(self):
        first = main.update_stock(self.master, self.folder([record()]), "first")
        before = digest(first.output_path)
        result = main.update_stock(first.output_path, self.folder([record()]), "first")
        self.assert_stopped(result, "継続")
        self.assertEqual(digest(first.output_path), before)
        self.assertEqual(self.outputs(), {first.output_path})

    def test_corrupt_workbook_stops(self):
        self.master.write_bytes(b"not a workbook")
        result = main.update_stock(self.master, self.folder([record()]), "continue")
        self.assert_stopped(result, "読み取れません")
        self.assertEqual(self.outputs(), set())

    def test_each_required_history_column_is_checked(self):
        first = main.update_stock(self.master, self.folder([record()]), "first")
        for name in main.HISTORY_HEADERS:
            with self.subTest(column=name):
                broken = self.base / f"broken_{main.HISTORY_HEADERS.index(name)}.xlsx"
                shutil.copyfile(first.output_path, broken)
                self.edit_master(lambda wb: setattr(wb["処理履歴"].cell(1, main.HISTORY_HEADERS.index(name) + 1), "value", "欠損"), broken)
                result = main.update_stock(broken, self.folder([record(slip="NEW")]), "continue")
                self.assert_stopped(result, "必須列")
                self.assertEqual(self.outputs(), {first.output_path})

    def test_history_duplicate_keys_identical_or_conflicting_are_fatal(self):
        first = main.update_stock(self.master, self.folder([record()]), "first")
        for conflict in (False, True):
            with self.subTest(conflict=conflict):
                broken = self.base / f"broken_{conflict}.xlsx"
                shutil.copyfile(first.output_path, broken)
                def change(wb):
                    ws = wb["処理履歴"]
                    values = [cell.value for cell in ws[2]]
                    if conflict:
                        values[5] += 1
                    ws.append(values)
                self.edit_master(change, broken)
                result = main.update_stock(broken, self.folder([record(slip="NEW")]), "continue")
                self.assert_stopped(result, "重複キー異常")
                self.assertEqual(self.outputs(), {first.output_path})

    def test_invalid_history_values_and_formulas_are_fatal(self):
        first = main.update_stock(self.master, self.folder([record()]), "first")
        for column, value in ((1, None), (2, ""), (3, "2026-02-30"), (4, 125), (5, "返品"), (6, 0), (7, "bad"), (6, "=1+1")):
            with self.subTest(column=column, value=value):
                broken = self.base / "broken.xlsx"
                shutil.copyfile(first.output_path, broken)
                self.edit_master(lambda wb: setattr(wb["処理履歴"].cell(2, column), "value", value), broken)
                result = main.update_stock(broken, self.folder([record(slip="NEW")]), "continue")
                self.assert_stopped(result, "処理履歴が不正")
                self.assertEqual(self.outputs(), {first.output_path})

    def test_history_header_only_is_valid_and_reordered_columns_are_supported(self):
        def change(wb):
            ws = wb.create_sheet("処理履歴")
            ws.append(list(reversed(main.HISTORY_HEADERS)))
        self.edit_master(change)
        result = main.update_stock(self.master, self.folder([record()]), "continue")
        self.assert_success(result)
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["処理履歴"]["D2"].value, "00125")
            self.assertEqual(wb["処理履歴"]["B2"].value, 5)

    def test_input_files_are_byte_identical_on_success_row_errors_and_fatal(self):
        for scenario in ("normal", "errors", "conflict"):
            with self.subTest(scenario=scenario):
                folder = self.sample_folder(scenario)
                paths = [self.master, *folder.glob("*.csv")]
                before = {path: digest(path) for path in paths}
                main.update_stock(self.master, folder, "first")
                self.assertEqual(before, {path: digest(path) for path in paths})

    def test_original_sheets_order_formulas_formatting_and_tables_are_preserved(self):
        result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(result)
        with closing(load_workbook(self.master)) as source, closing(load_workbook(result.output_path)) as output:
            self.assertEqual(output.sheetnames[:len(source.sheetnames)], source.sheetnames)
            for source_ws in source:
                output_ws = output[source_ws.title]
                self.assertEqual(output_ws.sheet_state, source_ws.sheet_state)
                self.assertEqual(str(output_ws.merged_cells), str(source_ws.merged_cells))
                self.assertEqual(output_ws.freeze_panes, source_ws.freeze_panes)
                self.assertEqual(output_ws.sheet_view.showGridLines, source_ws.sheet_view.showGridLines)
                self.assertEqual(set(output_ws.tables), set(source_ws.tables))
                for name, table in source_ws.tables.items():
                    self.assertEqual(output_ws.tables[name].ref, source_ws.tables[name].ref)
                for col, dimension in source_ws.column_dimensions.items():
                    self.assertEqual(output_ws.column_dimensions[col].width, dimension.width)
                for row, dimension in source_ws.row_dimensions.items():
                    self.assertEqual(output_ws.row_dimensions[row].height, dimension.height)
                for row in source_ws:
                    for cell in row:
                        output_cell = output_ws[cell.coordinate]
                        if not (source_ws.title == "商品マスタ" and cell.coordinate == "C2"):
                            self.assertEqual(output_cell.value, cell.value, (source_ws.title, cell.coordinate))
                        for attribute in ("font", "fill", "border", "alignment", "number_format", "protection"):
                            self.assertEqual(copy(getattr(output_cell, attribute)), copy(getattr(cell, attribute)), (source_ws.title, cell.coordinate, attribute))
                        self.assertEqual(output_cell.data_type, cell.data_type)

    def test_existing_report_sheets_and_their_formulas_are_not_modified(self):
        def change(wb):
            for title in ("集計", "エラー一覧", "警告一覧"):
                ws = wb.create_sheet(title)
                ws["A1"] = "既存データ"
                ws["A2"] = "=1+2"
        self.edit_master(change)
        first = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(first)
        for name in ("集計", "エラー一覧", "警告一覧"):
            self.assertEqual(first.report_sheets[name], name + "_2")
        with closing(load_workbook(first.output_path)) as wb:
            for name in ("集計", "エラー一覧", "警告一覧"):
                self.assertEqual(wb[name]["A2"].value, "=1+2")

    def test_filename_collision_uses_counter_and_never_overwrites(self):
        existing = self.base / "商品マスタ_更新結果_20261027_153000.xlsx"
        existing.write_bytes(b"existing result")
        before = digest(existing)
        with mock.patch.object(main, "datetime", FixedDateTime):
            first = main.update_stock(self.master, self.folder([record()]), "first")
            second = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(first)
        self.assert_success(second)
        self.assertEqual(first.output_path.name, "商品マスタ_更新結果_20261027_153000_001.xlsx")
        self.assertEqual(second.output_path.name, "商品マスタ_更新結果_20261027_153000_002.xlsx")
        self.assertEqual(digest(existing), before)

    @unittest.skipIf(os.name == "nt", "POSIXの排他的公開競合テスト")
    def test_filename_race_is_retried_without_overwriting_competing_file(self):
        real_link = os.link
        competitors = []
        def race(source, target):
            if not competitors:
                Path(target).write_bytes(b"another writer")
                competitors.append(Path(target))
            return real_link(source, target)
        with mock.patch.object(main.os, "link", side_effect=race):
            result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(result)
        self.assertEqual(competitors[0].read_bytes(), b"another writer")
        self.assertNotEqual(result.output_path, competitors[0])
        self.assertEqual(list(self.base.glob(".inventory-*.tmp")), [])

    def test_single_save_contains_complete_workbook_and_history_before_publish(self):
        original_save = Workbook.save
        real_publish = main._publish
        saved = []
        def save(workbook, destination):
            saved.append(destination)
            return original_save(workbook, destination)
        def inspect(payload, directory):
            with closing(load_workbook(io.BytesIO(payload))) as wb:
                self.assertEqual(wb["商品マスタ"]["C2"].value, 105)
                self.assertEqual(wb["処理履歴"].max_row, 2)
                for title in ("集計", "エラー一覧", "警告一覧", "処理履歴"):
                    self.assertIn(title, wb.sheetnames)
            return real_publish(payload, directory)
        with mock.patch.object(Workbook, "save", new=save), mock.patch.object(main, "_publish", side_effect=inspect):
            result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_success(result)
        self.assertEqual(len(saved), 1)
        self.assertIsInstance(saved[0], io.BytesIO)

    def test_workbook_serialization_failure_does_not_publish_or_register(self):
        before = digest(self.master)
        with mock.patch.object(Workbook, "save", side_effect=OSError("serialization failure")):
            result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_stopped(result, "serialization failure")
        self.assertEqual(self.outputs(), set())
        self.assertEqual(digest(self.master), before)
        self.assertEqual(list(self.base.glob(".inventory-*.tmp")), [])

    def test_failed_disk_write_leaves_no_result_and_retry_is_not_deduplicated(self):
        folder = self.folder([record()])
        before = digest(self.master)
        with mock.patch.object(main.os, "fsync", side_effect=OSError("disk full")):
            failed = main.update_stock(self.master, folder, "first")
        self.assert_stopped(failed, "disk full")
        self.assertEqual(self.outputs(), set())
        self.assertEqual(list(self.base.glob(".inventory-*.tmp")), [])
        self.assertEqual(digest(self.master), before)
        retried = main.update_stock(self.master, folder, "first")
        self.assert_success(retried)
        self.assertEqual((retried.applied_count, retried.duplicate_count), (1, 0))

    def test_failed_publish_cleans_temporary_file(self):
        method = "rename" if os.name == "nt" else "link"
        with mock.patch.object(main.os, method, side_effect=PermissionError("publish denied")):
            result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_stopped(result, "publish denied")
        self.assertEqual(self.outputs(), set())
        self.assertEqual(list(self.base.glob(".inventory-*.tmp")), [])

    def test_leftover_incomplete_temp_file_is_explicit_error(self):
        with mock.patch.object(main.os, "fsync", side_effect=OSError("disk full")), mock.patch.object(Path, "unlink", side_effect=PermissionError("cleanup denied")):
            result = main.update_stock(self.master, self.folder([record()]), "first")
        self.assert_stopped(result, "一時ファイルが残っています")
        leftovers = list(self.base.glob(".inventory-*.tmp"))
        self.assertEqual(len(leftovers), 1)
        self.assertTrue(any(str(leftovers[0]) in issue.message for issue in result.errors))
        self.assertEqual(self.outputs(), set())

    def test_missing_or_ambiguous_csv_headers_stop_entire_run(self):
        for headers in ([name for name in HEADERS if name != "明細ID"], HEADERS + ["日付"], HEADERS + ["日時"]):
            with self.subTest(headers=headers):
                folder = self.folder([record()])
                self.write_csv(folder / "later.csv", [record()], headers)
                result = main.update_stock(self.master, folder, "first")
                self.assert_stopped(result, "CSV必須列")
                self.assertEqual(self.outputs(), set())

    def test_bad_csv_encoding_and_malformed_quotes_stop(self):
        for payload in (b"\xff\xff\xff", ','.join(HEADERS).encode() + b'\n"unterminated'):
            with self.subTest(payload=payload):
                folder = self.folder([record()])
                (folder / "later.csv").write_bytes(payload)
                result = main.update_stock(self.master, folder, "first")
                self.assert_stopped(result)
                self.assertEqual(self.outputs(), set())

    def test_empty_csv_is_fatal_but_header_only_is_valid(self):
        folder = self.folder()
        (folder / "empty.csv").write_bytes(b"")
        self.assert_stopped(main.update_stock(self.master, folder, "first"), "CSVが空")
        result = main.update_stock(self.master, self.folder([]), "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 0)
        self.assertEqual(len(self.history(result.output_path)), 1)

    def test_missing_folder_or_no_csv_stops(self):
        for folder in (self.base / "missing", self.folder()):
            with self.subTest(folder=folder):
                self.assert_stopped(main.update_stock(self.master, folder, "first"))
                self.assertEqual(self.outputs(), set())

    def test_csv_datetime_alias_reordered_headers_and_uppercase_extension(self):
        folder = self.folder()
        headers = ["数量", "区分", "商品コード", "日時", "明細ID", "伝票番号", "任意列"]
        self.write_csv(folder / "input.CSV", [["5", "入庫", "00125", "2026-10-01 12:34:56", "001", "D001", "そのまま無視"]], headers, encoding="utf-8")
        nested = folder / "nested"
        nested.mkdir()
        self.write_csv(nested / "ignored.csv", [record(quantity="100")])
        result = main.update_stock(self.master, folder, "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        self.assertEqual(self.history(result.output_path)[1][2], "2026-10-01 12:34:56")

    def test_date_representation_whitespace_and_integer_representation_normalize(self):
        rows = [record(), record(day="2026/10/01", quantity="005"), record(day="2026-10-01 00:00:00", code=" 00125 ")]
        result = main.update_stock(self.master, self.folder(rows), "first")
        self.assert_success(result)
        self.assertEqual((result.applied_count, result.duplicate_count), (1, 2))

    def test_time_difference_is_conflict_and_invalid_dates_are_row_errors(self):
        result = main.update_stock(self.master, self.folder([record(day="2026-10-01 12:00"), record(day="2026-10-01 13:00")]), "first")
        self.assert_stopped(result, "矛盾")
        for day in ("2026-02-30", "10/01/2026", "2026-10-01T12:00:00+09:00"):
            with self.subTest(day=day):
                result = main.update_stock(self.master, self.folder([record(day=day)]), "first")
                self.assert_success(result)
                self.assertEqual(result.applied_count, 0)
                self.assertEqual(len(result.errors), 1)

    def test_non_positive_non_integer_quantities_are_row_errors(self):
        for quantity in ("0", "-1", "1.5", "1.0", "+1", "abc", "１", ""):
            with self.subTest(quantity=quantity):
                result = main.update_stock(self.master, self.folder([record(quantity=quantity)]), "first")
                self.assert_success(result)
                self.assertEqual(result.applied_count, 0)
                self.assertEqual(len(result.errors), 1)
                self.assertEqual(len(self.history(result.output_path)), 1)

    def test_multiline_csv_reports_physical_start_line_and_skips_blank_records(self):
        folder = self.folder([record(slip="MULTI\nLINE"), record(slip="BAD", detail="")])
        with (folder / "data.csv").open("a", encoding="utf-8") as stream:
            stream.write("\n,,,,,\n")
        result = main.update_stock(self.master, folder, "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        self.assertEqual(len(result.errors), 1)
        self.assertEqual(result.errors[0].row, 4)

    def test_bad_shape_row_does_not_suppress_valid_identical_key(self):
        folder = self.folder([record() + ["extra"], record()])
        result = main.update_stock(self.master, folder, "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 1)
        self.assertEqual(result.duplicate_count, 0)
        self.assertEqual(len(result.errors), 1)
        self.assertEqual(len(self.history(result.output_path)), 2)

    def test_csv_text_is_saved_as_literal_not_as_formula(self):
        result = main.update_stock(self.master, self.folder([record(slip="=1+1", detail="=2+2")]), "first")
        self.assert_success(result)
        with closing(load_workbook(result.output_path)) as wb:
            self.assertEqual(wb["処理履歴"]["A2"].value, "=1+1")
            self.assertEqual(wb["処理履歴"]["A2"].data_type, "s")
            self.assertEqual(wb["処理履歴"]["B2"].data_type, "s")

    def test_scale_2000_products_and_10_csv_files_of_1000_rows(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "商品マスタ"
        ws.append(["商品コード", "商品名", "現在庫", "単価"])
        for number in range(2000):
            ws.append([f"{number:05d}", f"商品{number}", 100, 50])
        wb.save(self.master)
        wb.close()
        folder = self.folder()
        for file_index in range(10):
            rows = [record(slip=f"S{file_index:02d}", detail=f"{row:04d}", code=f"{(file_index * 1000 + row) % 2000:05d}", quantity="1") for row in range(1000)]
            self.write_csv(folder / f"{file_index:02d}.csv", rows)
        result = main.update_stock(self.master, folder, "first")
        self.assert_success(result)
        self.assertEqual(result.applied_count, 10000)
        self.assertEqual(len(self.history(result.output_path)), 10001)
        self.assertEqual(set(self.stocks(result.output_path).values()), {105})


if __name__ == "__main__":
    unittest.main(verbosity=2)
