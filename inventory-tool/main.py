"""商品マスタと入出庫 CSV から、入力を変更せず更新結果 Excel を作成する。"""

from __future__ import annotations

import csv
import io
import math
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


MASTER_SHEET = "商品マスタ"
HISTORY_SHEET = "処理履歴"
MASTER_HEADERS = ("商品コード", "商品名", "現在庫")
CSV_HEADERS = ("伝票番号", "明細ID", "商品コード", "区分", "数量")
HISTORY_HEADERS = ("伝票番号", "明細ID", "日付", "商品コード", "区分", "数量", "処理日時")
SUMMARY_HEADERS = ("商品コード", "商品名", "現在庫", "今回入庫", "今回出庫", "処理後在庫", "備考")
ISSUE_HEADERS = ("元ファイル名", "行番号", "内容")


@dataclass(frozen=True)
class Issue:
    source_file: str
    row: int
    message: str


@dataclass
class RunResult:
    success: bool = False
    output_path: Path | None = None
    errors: list[Issue] = field(default_factory=list)
    warnings: list[Issue] = field(default_factory=list)
    applied_count: int = 0
    duplicate_count: int = 0
    report_sheets: dict[str, str] = field(default_factory=dict)


@dataclass
class Product:
    code: str
    name: str
    row: int
    stock: Any
    formula: bool
    invalid: bool = False


@dataclass(frozen=True)
class Transaction:
    slip: str
    detail: str
    day: str
    code: str
    kind: str
    quantity: int
    source_file: str
    row: int

    @property
    def key(self) -> tuple[str, str]:
        return self.slip, self.detail


class StopRun(Exception):
    """入力全体を安全に処理できない。結果ファイルを作らず停止する。"""


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _day(value: Any) -> str:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, date):
        return value.isoformat()
    else:
        raw = _text(value).replace("/", "-")
        if not re.fullmatch(
            r"\d{4}-\d{2}-\d{2}(?:[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d{1,6})?)?)?",
            raw,
            flags=re.ASCII,
        ):
            raise ValueError("日付の形式が不正")
        parsed = datetime.fromisoformat(raw)
    if parsed.tzinfo is not None:
        raise ValueError("日付にタイムゾーンは指定不可")
    if parsed.time() == time():
        return parsed.date().isoformat()
    return parsed.isoformat(sep=" ")


def _quantity(value: Any) -> int:
    if isinstance(value, bool):
        raise ValueError("数量は正の整数")
    if isinstance(value, int):
        number = value
    elif isinstance(value, str) and re.fullmatch(r"[0-9]+", value.strip()):
        number = int(value.strip())
    else:
        raise ValueError("数量は正の整数")
    if number <= 0:
        raise ValueError("数量は正の整数")
    return number


def _headers(ws: Any, required: tuple[str, ...], source: str, errors: list[Issue]) -> dict[str, int]:
    names = [_text(cell.value) for cell in ws[1]]
    invalid = [name for name in required if names.count(name) != 1]
    if invalid:
        errors.append(Issue(source, 1, f"{ws.title}: 必須列の不足または重複: {', '.join(invalid)}"))
        raise StopRun
    return {name: names.index(name) + 1 for name in required}


def _read_master(wb: Any, source: str, errors: list[Issue]) -> tuple[dict[str, Product], dict[str, int]]:
    if MASTER_SHEET not in wb.sheetnames:
        errors.append(Issue(source, 0, "シート「商品マスタ」がありません"))
        raise StopRun
    ws = wb[MASTER_SHEET]
    columns = _headers(ws, MASTER_HEADERS, source, errors)
    products: dict[str, Product] = {}
    duplicate = False
    for row in ws.iter_rows(min_row=2):
        if all(cell.value is None or cell.value == "" for cell in row):
            continue
        row_number = row[0].row
        code_cell = row[columns["商品コード"] - 1]
        code = _text(code_cell.value)
        if not code or code_cell.data_type == "f":
            errors.append(Issue(source, row_number, "商品マスタ: 商品コードが空欄または数式"))
            continue
        if code in products:
            errors.append(Issue(source, row_number, f"商品マスタ: 商品コード重複「{code}」（先行行 {products[code].row}）"))
            duplicate = True
            continue
        name_cell = row[columns["商品名"] - 1]
        name = _text(name_cell.value)
        stock_cell = row[columns["現在庫"] - 1]
        formula = stock_cell.data_type == "f"
        problems = []
        if not name:
            problems.append("商品名が空欄")
        stock = stock_cell.value
        if not formula and (
            isinstance(stock, bool)
            or not isinstance(stock, (int, float))
            or not math.isfinite(stock)
        ):
            problems.append("現在庫は数値セルである必要があります")
        if problems:
            errors.append(Issue(source, row_number, "商品マスタ: " + " / ".join(problems)))
        products[code] = Product(code, name, row_number, stock, formula, bool(problems))
    if duplicate:
        raise StopRun
    return products, columns


def _signature(code: Any, kind: Any, quantity: Any, day: Any) -> tuple[Any, ...]:
    # 不正行もキー比較の対象。正常行との食い違いを行エラーとして見逃さない。
    try:
        normalized_quantity = ("integer", _quantity(quantity))
    except (ValueError, OverflowError):
        normalized_quantity = ("invalid", _text(quantity))
    try:
        normalized_day = ("date", _day(day))
    except (ValueError, TypeError, OverflowError):
        normalized_day = ("invalid", _text(day))
    return _text(code), _text(kind), normalized_quantity, normalized_day


def _read_history(wb: Any, mode: str, source: str, errors: list[Issue]) -> tuple[dict, dict[str, int]]:
    if mode == "first":
        if HISTORY_SHEET in wb.sheetnames:
            errors.append(Issue(source, 1, "初回モードでは処理履歴のないオリジナルマスタを指定してください。履歴のある結果には継続を選択してください"))
            raise StopRun
        return {}, {}
    if HISTORY_SHEET not in wb.sheetnames:
        errors.append(Issue(source, 0, "継続モードには処理履歴が必要です。自動で初回扱いにはしません"))
        raise StopRun
    ws = wb[HISTORY_SHEET]
    columns = _headers(ws, HISTORY_HEADERS, source, errors)
    seen: dict = {}
    invalid = False
    for cells in ws.iter_rows(min_row=2):
        if all(cell.value is None or cell.value == "" for cell in cells):
            continue
        row = cells[0].row
        values = {name: cells[col - 1].value for name, col in columns.items()}
        try:
            if any(cells[col - 1].data_type == "f" for col in columns.values()):
                raise ValueError("必須項目に数式があります")
            if any(not _text(value) for value in values.values()):
                raise ValueError("必須項目が空欄")
            if not isinstance(values["商品コード"], str):
                raise ValueError("商品コードは文字列である必要があります")
            if _text(values["区分"]) not in ("入庫", "出庫"):
                raise ValueError("区分が不正")
            _quantity(values["数量"])
            _day(values["日付"])
            processed = values["処理日時"]
            if not isinstance(processed, datetime):
                datetime.fromisoformat(_text(processed))
            key = _text(values["伝票番号"]), _text(values["明細ID"])
            if key in seen:
                raise ValueError(f"重複キー異常（先行行 {seen[key][2]}）")
            seen[key] = (
                _signature(values["商品コード"], values["区分"], values["数量"], values["日付"]),
                source,
                row,
            )
        except (ValueError, TypeError, OverflowError) as exc:
            errors.append(Issue(source, row, f"処理履歴が不正: {exc}"))
            invalid = True
    if invalid:
        raise StopRun
    return seen, columns


def _csv_records(path: Path, errors: list[Issue]):
    try:
        data = path.read_bytes()
        try:
            content = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = data.decode("cp932")
    except (OSError, UnicodeError) as exc:
        errors.append(Issue(path.name, 0, f"CSVを読めません（UTF-8 / CP932）: {exc}"))
        raise StopRun from exc
    reader = csv.reader(io.StringIO(content, newline=""), strict=True)
    try:
        names = [_text(value) for value in next(reader)]
        date_names = [name for name in ("日付", "日時") if name in names]
        invalid = [name for name in CSV_HEADERS if names.count(name) != 1]
        if len(date_names) != 1 or names.count(date_names[0]) != 1:
            invalid.append("日付（または日時のどちらか1列）")
        if invalid:
            errors.append(Issue(path.name, 1, "CSV必須列の不足または重複: " + ", ".join(invalid)))
            raise StopRun
        day_name = date_names[0]
        positions = {name: names.index(name) for name in (*CSV_HEADERS, day_name)}
        while True:
            line = reader.line_num + 1
            try:
                record = next(reader)
            except StopIteration:
                break
            if not record or all(not value.strip() for value in record):
                continue
            fields = {name: _text(record[pos]) if pos < len(record) else "" for name, pos in positions.items()}
            fields["日付"] = fields[day_name]
            yield line, fields, len(record) == len(names)
    except StopIteration as exc:
        errors.append(Issue(path.name, 1, "CSVが空です（見出しなし）"))
        raise StopRun from exc
    except csv.Error as exc:
        errors.append(Issue(path.name, reader.line_num or 1, f"CSV構造が不正: {exc}"))
        raise StopRun from exc


def _read_transactions(folder: Path, history: dict, result: RunResult) -> list[Transaction]:
    try:
        paths = sorted((p for p in folder.iterdir() if p.is_file() and p.suffix.lower() == ".csv"), key=lambda p: p.name)
    except OSError as exc:
        result.errors.append(Issue(folder.name, 0, f"CSVフォルダを読めません: {exc}"))
        raise StopRun from exc
    if not paths:
        result.errors.append(Issue(folder.name, 0, "指定フォルダ直下にCSVファイルがありません"))
        raise StopRun
    seen = dict(history)
    accepted_keys: set[tuple[str, str]] = set()
    transactions = []
    conflict = False
    for path in paths:
        for row, fields, shape_ok in _csv_records(path, result.errors):
            key = fields["伝票番号"], fields["明細ID"]
            signature = _signature(fields["商品コード"], fields["区分"], fields["数量"], fields["日付"])
            duplicate = False
            if all(key):
                if key in seen:
                    previous, source, previous_row = seen[key]
                    if signature != previous:
                        result.errors.append(Issue(path.name, row, f"同一キーの内容矛盾: {key[0]} / {key[1]}（照合先 {source}, 行 {previous_row}）。実行全体を停止"))
                        conflict = True
                    else:
                        duplicate = key in history or key in accepted_keys
                else:
                    seen[key] = signature, path.name, row
            problems = []
            if not shape_ok:
                problems.append("列数が見出しと一致しません")
            for name in (*CSV_HEADERS, "日付"):
                if not fields[name]:
                    problems.append(f"{name}が空欄")
            if fields["区分"] and fields["区分"] not in ("入庫", "出庫"):
                problems.append("区分は入庫または出庫")
            try:
                quantity = _quantity(fields["数量"])
            except (ValueError, OverflowError):
                quantity = 0
                if fields["数量"]:
                    problems.append("数量は正の整数")
            try:
                day = _day(fields["日付"])
            except (ValueError, TypeError, OverflowError):
                day = ""
                if fields["日付"]:
                    problems.append("日付が不正")
            if problems:
                result.errors.append(Issue(path.name, row, " / ".join(problems)))
            elif duplicate:
                result.duplicate_count += 1
            else:
                transactions.append(Transaction(*key, day, fields["商品コード"], fields["区分"], quantity, path.name, row))
                accepted_keys.add(key)
    if conflict:
        raise StopRun
    return transactions


def _set_cell(ws: Any, row: int, column: int, value: Any) -> None:
    cell = ws.cell(row, column)
    cell.value = value
    if isinstance(value, str):
        # CSV由来の「=...」等も数式に変換せず、そのまま文字列として保存。
        cell.data_type = "s"


def _new_report(wb: Any, title: str, headers: tuple[str, ...]) -> Any:
    name = title
    counter = 2
    while name in wb.sheetnames:
        name = f"{title}_{counter}"
        counter += 1
    ws = wb.create_sheet(name)
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
        ws.column_dimensions[cell.column_letter].width = 22 if header != "内容" else 85
    ws.freeze_panes = "A2"
    return ws


def _build_result(wb: Any, products: dict[str, Product], master_columns: dict[str, int],
                  transactions: list[Transaction], history_columns: dict[str, int],
                  source: str, result: RunResult) -> int:
    totals = {code: [0, 0] for code in products}
    applied = []
    sources: dict[str, list[Transaction]] = {}
    formula_warned = set()
    for transaction in transactions:
        product = products.get(transaction.code)
        if product is None:
            result.errors.append(Issue(transaction.source_file, transaction.row, f"未登録商品コード「{transaction.code}」"))
        elif product.invalid:
            result.errors.append(Issue(transaction.source_file, transaction.row, f"商品コード「{transaction.code}」はマスタ不正のため未更新（{source}, 行 {product.row}）"))
        elif product.formula:
            result.warnings.append(Issue(transaction.source_file, transaction.row, f"商品コード「{transaction.code}」: 現在庫が数式のため未更新（{source}, 行 {product.row}）"))
            formula_warned.add(product.code)
        else:
            totals[product.code][0 if transaction.kind == "入庫" else 1] += transaction.quantity
            applied.append(transaction)
            sources.setdefault(product.code, []).append(transaction)

    # ここまで在庫は未変更。全ファイル・全履歴の検証完了後だけメモリ上で反映。
    master = wb[MASTER_SHEET]
    summary_rows = []
    for code, product in products.items():
        incoming, outgoing = totals[code]
        after = product.stock
        note = ""
        _set_cell(master, product.row, master_columns["商品コード"], code)
        if product.invalid:
            note = "マスタ不正のため未更新"
        elif product.formula:
            note = "現在庫が数式のため未更新"
            if code not in formula_warned:
                result.warnings.append(Issue(source, product.row, f"商品マスタ: 商品コード「{code}」: {note}"))
        else:
            after = product.stock + incoming - outgoing
            if not math.isfinite(after):
                result.errors.append(Issue(source, product.row, f"商品コード「{code}」の計算結果がExcel数値範囲を超えました"))
                raise StopRun
            master.cell(product.row, master_columns["現在庫"]).value = after
            if after < 0:
                note = "負在庫（自動補正なし）"
                locations = sources.get(code) or [Issue(source, product.row, "")]
                for location in locations:
                    result.warnings.append(Issue(location.source_file, location.row, f"商品コード「{code}」: 処理後在庫 {after}（負在庫・自動補正なし）"))
        summary_rows.append((code, product.name, product.stock, incoming, outgoing, after, note))

    result.errors.sort(key=lambda issue: (issue.source_file, issue.row, issue.message))
    result.warnings.sort(key=lambda issue: (issue.source_file, issue.row, issue.message))
    summary = _new_report(wb, "集計", SUMMARY_HEADERS)
    result.report_sheets["集計"] = summary.title
    for row, values in enumerate(summary_rows, 2):
        for col, value in enumerate(values, 1):
            _set_cell(summary, row, col, value)
        summary.cell(row, 1).number_format = "@"
    for title, issues in (("エラー一覧", result.errors), ("警告一覧", result.warnings)):
        ws = _new_report(wb, title, ISSUE_HEADERS)
        result.report_sheets[title] = ws.title
        for row, issue in enumerate(issues, 2):
            for col, value in enumerate((issue.source_file, issue.row, issue.message), 1):
                _set_cell(ws, row, col, value)

    if HISTORY_SHEET in wb.sheetnames:
        history = wb[HISTORY_SHEET]
    else:
        history = _new_report(wb, HISTORY_SHEET, HISTORY_HEADERS)
        history_columns = {name: index for index, name in enumerate(HISTORY_HEADERS, 1)}
    processed_at = datetime.now().astimezone().isoformat(sep=" ", timespec="seconds")
    row = history.max_row + 1
    for transaction in applied:
        values = (transaction.slip, transaction.detail, transaction.day, transaction.code,
                  transaction.kind, transaction.quantity, processed_at)
        for name, value in zip(HISTORY_HEADERS, values):
            _set_cell(history, row, history_columns[name], value)
        history.cell(row, history_columns["商品コード"]).number_format = "@"
        row += 1
    return len(applied)


def _publish(payload: bytes, directory: Path) -> Path:
    """完成バイト列を同一フォルダの一時ファイル経由で原子的・排他的に公開。"""
    temporary: Path | None = None
    published: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(mode="wb", dir=directory, prefix=".inventory-", suffix=".tmp", delete=False) as stream:
            temporary = Path(stream.name)
            stream.write(payload)
            stream.flush()
            os.fsync(stream.fileno())
        stem = "商品マスタ_更新結果_" + datetime.now().strftime("%Y%m%d_%H%M%S")
        counter = 0
        while True:
            suffix = "" if counter == 0 else f"_{counter:03d}"
            target = directory / f"{stem}{suffix}.xlsx"
            try:
                if os.name == "nt":
                    # Windowsのrenameは既存宛先を上書きしない。
                    os.rename(temporary, target)
                else:
                    # POSIX renameは上書きするため、排他的なhard linkで公開。
                    os.link(temporary, target)
                published = target
                break
            except FileExistsError:
                counter += 1
        if temporary.exists():
            temporary.unlink()
        return published
    except Exception as exc:
        remaining = ""
        if temporary is not None and temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                remaining = f"。一時ファイルが残っています（未完了扱い）: {temporary}"
        if published is not None:
            remaining += f"。完成ファイルはありますが後処理に失敗しました（未完了扱い）: {published}"
        raise OSError(f"結果保存に失敗: {exc}{remaining}") from exc


def update_stock(master_path: str | Path, csv_folder: str | Path, mode: str) -> RunResult:
    """modeは明示的に 'first' または 'continue'。保存完了時だけsuccess=True。"""
    result = RunResult()
    master_path, csv_folder = Path(master_path), Path(csv_folder)
    wb = None
    try:
        if mode not in ("first", "continue"):
            result.errors.append(Issue(master_path.name, 0, "初回 / 継続モードを明示選択してください"))
            raise StopRun
        if master_path.suffix.lower() != ".xlsx":
            result.errors.append(Issue(master_path.name, 0, "入力マスタは.xlsx形式を指定してください"))
            raise StopRun
        try:
            wb = load_workbook(master_path, data_only=False, keep_links=True, rich_text=True)
        except Exception as exc:
            result.errors.append(Issue(master_path.name, 0, f"マスタ / 処理履歴を読み取れません: {exc}"))
            raise StopRun from exc
        products, master_columns = _read_master(wb, master_path.name, result.errors)
        history, history_columns = _read_history(wb, mode, master_path.name, result.errors)
        transactions = _read_transactions(csv_folder, history, result)
        applied_count = _build_result(wb, products, master_columns, transactions, history_columns, master_path.name, result)
        with io.BytesIO() as buffer:
            # 処理履歴を含む完成Workbookを一度だけ保存。入力ファイルには一切保存しない。
            wb.save(buffer)
            payload = buffer.getvalue()
        result.output_path = _publish(payload, master_path.parent)
        result.applied_count = applied_count
        result.success = True
    except StopRun:
        pass
    except Exception as exc:
        result.errors.append(Issue(master_path.name, 0, f"実行失敗（処理完了・履歴登録とは扱いません）: {exc}"))
    finally:
        if wb is not None:
            wb.close()
    return result


def main() -> None:
    import tkinter as tk
    from tkinter import filedialog, messagebox, scrolledtext, ttk

    root = tk.Tk()
    root.title("商品マスタ 在庫更新ツール")
    root.columnconfigure(1, weight=1)
    mode = tk.StringVar(value="")
    master = tk.StringVar()
    folder = tk.StringVar()
    ttk.Label(root, text="1. 初回 / 継続を選択してください").grid(row=0, column=0, columnspan=3, sticky="w", padx=12, pady=10)
    choices = ttk.Frame(root)
    choices.grid(row=1, column=0, columnspan=3, sticky="w", padx=12)
    ttk.Radiobutton(choices, text="初回（履歴のないオリジナル）", variable=mode, value="first").pack(side="left")
    ttk.Radiobutton(choices, text="継続（前回の更新結果）", variable=mode, value="continue").pack(side="left")

    def choose_master() -> None:
        if not mode.get():
            messagebox.showerror("選択が必要です", "先に初回 / 継続を選択してください", parent=root)
            return
        selected = filedialog.askopenfilename(parent=root, title="商品マスタ Excel を指定", filetypes=[("Excel", "*.xlsx")])
        if selected:
            master.set(selected)

    def choose_folder() -> None:
        if not mode.get() or not master.get():
            messagebox.showerror("選択が必要です", "初回 / 継続、マスタの順に指定してください", parent=root)
            return
        selected = filedialog.askdirectory(parent=root, title="入出庫 CSV フォルダを指定")
        if selected:
            folder.set(selected)

    for row, label, variable, command in (
        (2, "2. マスタ Excel", master, choose_master),
        (3, "3. CSV フォルダ", folder, choose_folder),
    ):
        ttk.Label(root, text=label).grid(row=row, column=0, sticky="w", padx=12, pady=10)
        ttk.Entry(root, textvariable=variable, width=65, state="readonly").grid(row=row, column=1, sticky="ew")
        ttk.Button(root, text="選択", command=command).grid(row=row, column=2, padx=12)

    def execute() -> None:
        if not mode.get() or not master.get() or not folder.get():
            messagebox.showerror("選択が必要です", "初回 / 継続、マスタ、CSVフォルダをすべて選択してください", parent=root)
            return
        button.configure(state="disabled")
        root.update_idletasks()
        try:
            result = update_stock(master.get(), folder.get(), mode.get())
        finally:
            button.configure(state="normal")
        window = tk.Toplevel(root)
        window.title("保存完了" if result.success else "処理停止")
        output = scrolledtext.ScrolledText(window, width=105, height=28, wrap="word")
        output.pack(fill="both", expand=True, padx=12, pady=12)
        if result.success:
            output.insert("end", f"保存完了: {result.output_path}\n反映 {result.applied_count} 行 / 同一内容の重複除外 {result.duplicate_count} 行\n")
            output.insert("end", "今回の一覧シート: " + " / ".join(result.report_sheets.values()) + "\n")
        else:
            output.insert("end", "処理停止。処理完了・履歴登録とは扱いません。入力ファイルは変更していません。\n")
        for title, issues in (("エラー一覧", result.errors), ("警告一覧", result.warnings)):
            output.insert("end", f"\n{title} ({len(issues)} 件)\n")
            for issue in issues:
                output.insert("end", f"{issue.source_file} / 行 {issue.row}: {issue.message}\n")
        output.configure(state="disabled")

    button = ttk.Button(root, text="4. 実行", command=execute)
    button.grid(row=4, column=0, columnspan=3, pady=14)
    ttk.Label(root, text="出力先: 入力マスタと同じフォルダ。次回は更新結果を継続モードで選択。", wraplength=720).grid(row=5, column=0, columnspan=3, padx=12, pady=(0, 12))
    root.mainloop()


if __name__ == "__main__":
    main()
